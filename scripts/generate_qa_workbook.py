#!/usr/bin/env python3
"""Generate the Vietnamese QA workbook for the English Center MVC project."""

from __future__ import annotations

from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "EnglishCenter_QA_TestCases.xlsx"

NAVY = "17365D"
BLUE = "2F75B5"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "E7E6E6"
GREEN = "C6E0B4"
RED = "F4CCCC"
YELLOW = "FFF2CC"
ORANGE = "FCE4D6"
WHITE = "FFFFFF"
THIN_GRAY = Side(style="thin", color="B7B7B7")


def requirement(req_id: str, module: str, description: str, priority: str, source: str):
    return [req_id, module, description, priority, source, "Baseline v1.0"]


REQUIREMENTS = [
    requirement("REQ-AUTH-001", "Auth", "Người dùng hoạt động đăng nhập được bằng đúng tên đăng nhập và mật khẩu.", "P0", "AccountController.Login"),
    requirement("REQ-AUTH-002", "Auth", "Hệ thống từ chối tài khoản không tồn tại, bị khóa hoặc sai mật khẩu.", "P0", "AccountController.Login"),
    requirement("REQ-AUTH-003", "Auth", "Tên đăng nhập và mật khẩu là dữ liệu bắt buộc.", "P1", "LoginViewModel"),
    requirement("REQ-AUTH-004", "Auth", "Hỗ trợ remember-me và chỉ chuyển tiếp đến ReturnUrl nội bộ.", "P1", "AccountController.Login"),
    requirement("REQ-AUTH-005", "Auth", "Đăng xuất phải hủy phiên xác thực và quay về trang đăng nhập.", "P0", "AccountController.Logout"),
    requirement("REQ-AUTH-006", "Auth", "Chỉ Admin và GiaoVu được truy cập các module quản trị.", "P0", "Authorize attributes"),
    requirement("REQ-COURSE-001", "Course", "Danh sách khóa học hỗ trợ tìm kiếm và sắp xếp bản ghi mới trước.", "P1", "CourseController.Index"),
    requirement("REQ-COURSE-002", "Course", "Danh sách khóa học phân trang 5 bản ghi mỗi trang.", "P2", "CourseController.Index"),
    requirement("REQ-COURSE-003", "Course", "Xem được chi tiết khóa học và các lớp liên kết.", "P1", "CourseController.Details"),
    requirement("REQ-COURSE-004", "Course", "Tạo khóa học hợp lệ; dùng ảnh mặc định khi không tải ảnh.", "P0", "CourseController.Create"),
    requirement("REQ-COURSE-005", "Course", "Tải lên và thay thế ảnh bìa khóa học.", "P1", "CourseController.Create/Edit"),
    requirement("REQ-COURSE-006", "Course", "Kiểm tra tên, học phí, tổng số buổi và mô tả theo giới hạn model.", "P0", "Course model"),
    requirement("REQ-COURSE-007", "Course", "Cập nhật khóa học và giữ ảnh cũ khi không chọn ảnh mới.", "P1", "CourseController.Edit"),
    requirement("REQ-COURSE-008", "Course", "Chỉ xóa khóa học chưa có lớp liên kết.", "P0", "CourseController.Delete"),
    requirement("REQ-CLASS-001", "Class", "Danh sách lớp hỗ trợ tìm theo mã, tên lớp hoặc giảng viên.", "P1", "ClassController.Index"),
    requirement("REQ-CLASS-002", "Class", "Danh sách lớp hỗ trợ lọc khóa học và phân trang 5 bản ghi.", "P2", "ClassController.Index"),
    requirement("REQ-CLASS-003", "Class", "Chi tiết lớp hiển thị khóa học và danh sách ghi danh.", "P1", "ClassController.Details"),
    requirement("REQ-CLASS-004", "Class", "Tạo lớp hợp lệ và chỉ chọn khóa học đang hoạt động.", "P0", "ClassController.Create"),
    requirement("REQ-CLASS-005", "Class", "Không cho phép tạo lớp có ClassCode trùng.", "P0", "ClassController.Create"),
    requirement("REQ-CLASS-006", "Class", "Kiểm tra trường bắt buộc, độ dài và sĩ số 1-100.", "P0", "Class model"),
    requirement("REQ-CLASS-007", "Class", "Cập nhật thông tin lớp hợp lệ.", "P1", "ClassController.Edit"),
    requirement("REQ-CLASS-008", "Class", "Chỉ xóa lớp chưa có học viên ghi danh.", "P0", "ClassController.Delete"),
    requirement("REQ-ENR-001", "Enrollment", "Danh sách ghi danh hỗ trợ tìm kiếm, lọc và phân trang 5 bản ghi.", "P1", "EnrollmentController.Index"),
    requirement("REQ-ENR-002", "Enrollment", "Chi tiết ghi danh hiển thị học viên, lớp, học phí, điểm danh và điểm.", "P1", "EnrollmentController.Details"),
    requirement("REQ-ENR-003", "Enrollment", "Tạo ghi danh hợp lệ với giá trị mặc định phù hợp.", "P0", "EnrollmentController.Create"),
    requirement("REQ-ENR-004", "Enrollment", "Khi chọn lớp, hệ thống lấy học phí chuẩn của khóa học.", "P1", "GetClassDefaultFee"),
    requirement("REQ-ENR-005", "Enrollment", "Không ghi danh trùng học viên trong cùng một lớp.", "P0", "EnrollmentController.Create/Edit"),
    requirement("REQ-ENR-006", "Enrollment", "Không ghi danh vượt quá sĩ số tối đa của lớp.", "P0", "EnrollmentController.Create"),
    requirement("REQ-ENR-007", "Enrollment", "Tự động tính DaNopDu hoặc ConNo từ ActualFee và PaidAmount.", "P0", "EnrollmentController.Create/Edit"),
    requirement("REQ-ENR-008", "Enrollment", "Học phí thực tế và số tiền đã nộp nằm trong khoảng 0-100.000.000.", "P0", "Enrollment model"),
    requirement("REQ-ENR-009", "Enrollment", "Cập nhật trạng thái học tập và thông tin ghi danh.", "P1", "EnrollmentController.Edit"),
    requirement("REQ-ENR-010", "Enrollment", "Chỉ xóa ghi danh chưa có điểm danh hoặc điểm số.", "P0", "EnrollmentController.Delete"),
    requirement("REQ-ATT-001", "Attendance", "Danh sách điểm danh hỗ trợ tìm kiếm, lọc và phân trang 10 bản ghi.", "P1", "AttendanceController.Index"),
    requirement("REQ-ATT-002", "Attendance", "Sổ điểm danh chỉ lấy học viên có trạng thái DangHoc.", "P0", "AttendanceController.TakeAttendance"),
    requirement("REQ-ATT-003", "Attendance", "Tạo mới hoặc cập nhật điểm danh theo lớp và buổi học.", "P0", "AttendanceController.SaveAttendance"),
    requirement("REQ-ATT-004", "Attendance", "Lớp, ngày và số buổi là bắt buộc; số buổi nằm trong khoảng 1-200.", "P0", "AttendanceSheetViewModel"),
    requirement("REQ-ATT-005", "Attendance", "Ghi chú điểm danh tối đa 255 ký tự.", "P1", "Attendance model"),
    requirement("REQ-ATT-006", "Attendance", "Có thể xóa một bản ghi điểm danh.", "P1", "AttendanceController.Delete"),
    requirement("REQ-GRADE-001", "Grade", "Danh sách điểm hỗ trợ tìm kiếm, lọc và phân trang 10 bản ghi.", "P1", "GradeController.Index"),
    requirement("REQ-GRADE-002", "Grade", "Sổ điểm chỉ lấy học viên có trạng thái DangHoc và nạp điểm đã có.", "P0", "GradeController.EnterGrades"),
    requirement("REQ-GRADE-003", "Grade", "Chỉ tạo bản ghi điểm mới khi có điểm hoặc nhận xét; có thể cập nhật bản ghi cũ.", "P0", "GradeController.SaveGrades"),
    requirement("REQ-GRADE-004", "Grade", "Điểm từng kỹ năng nằm trong khoảng 0-990.", "P0", "GradeSheetViewModel"),
    requirement("REQ-GRADE-005", "Grade", "Kỳ kiểm tra bắt buộc, tối đa 50 ký tự; nhận xét tối đa 500 ký tự.", "P1", "Grade model"),
    requirement("REQ-GRADE-006", "Grade", "TOEIC Overall bằng Listening cộng Reading.", "P0", "GradeController.CalculateOverall"),
    requirement("REQ-GRADE-007", "Grade", "IELTS Overall là trung bình kỹ năng đã nhập và làm tròn band .0/.5.", "P0", "GradeController.CalculateOverall"),
    requirement("REQ-GRADE-008", "Grade", "Khóa học khác tính trung bình kỹ năng và làm tròn một chữ số.", "P1", "GradeController.CalculateOverall"),
    requirement("REQ-GRADE-009", "Grade", "Có thể xóa một bản ghi điểm.", "P1", "GradeController.Delete"),
    requirement("REQ-COMMON-001", "Common", "Người chưa đăng nhập bị chuyển đến Login khi vào route quản trị.", "P0", "Cookie authentication"),
    requirement("REQ-COMMON-002", "Common", "Các thao tác thay đổi dữ liệu phải kiểm tra anti-forgery token.", "P0", "ValidateAntiForgeryToken"),
    requirement("REQ-COMMON-003", "Common", "ID rỗng/không tồn tại trả về Not Found phù hợp.", "P1", "CRUD controllers"),
    requirement("REQ-COMMON-004", "Common", "Ứng dụng build và khởi tạo seed data thành công.", "P0", "Program/DbInitializer"),
    requirement("REQ-COMMON-005", "Common", "Giao diện luồng chính dùng được trên trình duyệt mục tiêu và không vỡ bố cục.", "P2", "MVC views"),
]


CASES: list[dict[str, str]] = []
MODULE_COUNTERS: Counter[str] = Counter()


def add_case(
    module: str,
    req_id: str,
    title: str,
    priority: str,
    test_type: str,
    precondition: str,
    test_data: str,
    steps: str,
    expected: str,
    severity: str = "Medium",
):
    MODULE_COUNTERS[module] += 1
    CASES.append(
        {
            "id": f"TC-{module.upper()}-{MODULE_COUNTERS[module]:03d}",
            "module": module,
            "req": req_id,
            "title": title,
            "priority": priority,
            "type": test_type,
            "precondition": precondition,
            "data": test_data,
            "steps": steps,
            "expected": expected,
            "actual": "",
            "status": "Chưa chạy",
            "severity": severity,
            "tester": "",
            "date": "",
            "note": "",
        }
    )


ADMIN = "Database seed sẵn; chưa đăng nhập."
ADMIN_IN = "Đã đăng nhập bằng admin/admin@123."
DATA_READY = "Đã đăng nhập Admin; có dữ liệu seed và dữ liệu QA cần thiết."

# Auth: 11 cases
add_case("Auth", "REQ-AUTH-001", "Đăng nhập thành công bằng tài khoản Admin", "P0", "Smoke", ADMIN, "admin / admin@123", "1. Mở /Account/Login\n2. Nhập tài khoản\n3. Chọn Đăng nhập", "Đăng nhập thành công, chuyển đến Home và phiên mang role Admin.", "Critical")
add_case("Auth", "REQ-AUTH-001", "Đăng nhập thành công bằng tài khoản GiaoVu", "P0", "Functional", ADMIN, "giaovu01 / giaovu@123", "1. Mở Login\n2. Nhập tài khoản GiaoVu\n3. Đăng nhập", "Đăng nhập thành công; GiaoVu truy cập được các module quản trị.", "Critical")
add_case("Auth", "REQ-AUTH-001", "Đăng nhập thành công bằng tài khoản SinhVien", "P1", "Functional", ADMIN, "sv2026001 / 123456", "1. Mở Login\n2. Nhập tài khoản SinhVien\n3. Đăng nhập", "Đăng nhập thành công, chuyển đến Home và phiên mang role SinhVien.", "High")
add_case("Auth", "REQ-AUTH-002", "Từ chối mật khẩu sai", "P0", "Negative", ADMIN, "admin / wrong-password", "1. Mở Login\n2. Nhập mật khẩu sai\n3. Đăng nhập", "Không tạo phiên; hiển thị 'Mật khẩu không đúng'.", "Critical")
add_case("Auth", "REQ-AUTH-002", "Từ chối tài khoản không tồn tại hoặc bị khóa", "P0", "Negative", ADMIN, "notfound / 123456; tài khoản IsActive=false", "1. Thử user không tồn tại\n2. Đặt một user test IsActive=false và thử lại", "Cả hai lần đều không tạo phiên; hiển thị tài khoản không tồn tại hoặc đã bị khóa.", "Critical")
add_case("Auth", "REQ-AUTH-003", "Kiểm tra trường đăng nhập bắt buộc", "P1", "Validation", ADMIN, "Username/password rỗng theo từng tổ hợp", "1. Để trống cả hai trường và submit\n2. Chỉ nhập username\n3. Chỉ nhập password", "Hiển thị validation đúng trường; không gửi đăng nhập thành công.", "High")
add_case("Auth", "REQ-AUTH-004", "Remember me tạo phiên đăng nhập kéo dài", "P1", "Functional", ADMIN, "admin / admin@123; RememberMe=true", "1. Chọn Ghi nhớ đăng nhập\n2. Đăng nhập\n3. Kiểm tra cookie xác thực", "Cookie có thuộc tính persistent và hạn dùng xấp xỉ 7 ngày; vẫn HttpOnly.", "Medium")
add_case("Auth", "REQ-AUTH-004", "Chuyển về ReturnUrl nội bộ sau đăng nhập", "P1", "Functional", ADMIN, "ReturnUrl=/Course", "1. Mở /Account/Login?returnUrl=/Course\n2. Đăng nhập Admin", "Chuyển đến /Course sau khi đăng nhập.", "High")
add_case("Auth", "REQ-AUTH-004", "Không chuyển hướng đến ReturnUrl bên ngoài", "P0", "Security", ADMIN, "ReturnUrl=https://example.com", "1. Mở Login với ReturnUrl bên ngoài\n2. Đăng nhập Admin", "Không redirect ra domain ngoài; chuyển đến Home.", "Critical")
add_case("Auth", "REQ-AUTH-005", "Đăng xuất hủy phiên", "P0", "Smoke", ADMIN_IN, "N/A", "1. Chọn Đăng xuất\n2. Truy cập lại /Course", "Quay về Login; truy cập /Course yêu cầu đăng nhập lại.", "Critical")
add_case("Auth", "REQ-AUTH-006", "SinhVien bị chặn khỏi trang quản trị", "P0", "Authorization", "Đã đăng nhập sv2026001/123456.", "/Course, /Class, /Enrollment, /Attendance, /Grade", "1. Mở lần lượt từng URL quản trị\n2. Ghi nhận response và trang đích", "Tất cả route bị từ chối và chuyển đến /Account/Denied; không lộ dữ liệu quản trị.", "Critical")

# Course: 14 cases
add_case("Course", "REQ-COURSE-001", "Hiển thị danh sách khóa học theo bản ghi mới nhất", "P1", "Smoke", ADMIN_IN, "Dữ liệu seed", "1. Mở /Course\n2. Đối chiếu tên và thứ tự bản ghi", "Danh sách hiển thị đúng dữ liệu; CourseId lớn hơn nằm trước.", "High")
add_case("Course", "REQ-COURSE-001", "Tìm khóa học theo một phần tên", "P1", "Functional", ADMIN_IN, "search=IELTS", "1. Nhập IELTS\n2. Tìm kiếm", "Chỉ trả các khóa học có CourseName chứa IELTS, không phân biệt phần còn lại của tên.", "Medium")
add_case("Course", "REQ-COURSE-001", "Tìm khóa học không có kết quả", "P2", "Negative", ADMIN_IN, "search=ZZZ-NOT-FOUND", "1. Nhập từ khóa\n2. Tìm kiếm", "Danh sách rỗng/hiển thị trạng thái không có dữ liệu; không phát sinh lỗi.", "Low")
add_case("Course", "REQ-COURSE-002", "Phân trang khóa học 5 bản ghi", "P2", "Functional", "Có ít nhất 7 khóa học; đã đăng nhập Admin.", "page=1 và page=2", "1. Mở trang 1\n2. Chuyển trang 2\n3. Quay lại trang 1", "Mỗi trang tối đa 5 bản ghi, không trùng/mất bản ghi; điều hướng đúng.", "Medium")
add_case("Course", "REQ-COURSE-003", "Xem chi tiết khóa học và lớp liên kết", "P1", "Functional", DATA_READY, "Khóa IELTS seed", "1. Mở Details của khóa IELTS\n2. Kiểm tra thông tin và danh sách lớp", "Thông tin khóa học đúng và lớp IELTS-K02 xuất hiện trong dữ liệu liên kết.", "High")
add_case("Course", "REQ-COURSE-004", "Tạo khóa học không upload ảnh", "P0", "Smoke", ADMIN_IN, "QA General English; 2500000; 20 buổi; mô tả hợp lệ", "1. Mở Create\n2. Nhập dữ liệu, bỏ trống ảnh\n3. Lưu", "Tạo thành công; ThumbnailUrl là /img/default-course.jpg; bản ghi xuất hiện trong danh sách.", "Critical")
add_case("Course", "REQ-COURSE-005", "Tạo khóa học với ảnh bìa hợp lệ", "P1", "Functional", ADMIN_IN, "Ảnh PNG/JPG nhỏ hơn 1 MB", "1. Nhập dữ liệu khóa học\n2. Chọn ảnh\n3. Lưu\n4. Mở chi tiết", "Tạo thành công; file có tên duy nhất trong /uploads/courses và ảnh hiển thị được.", "High")
add_case("Course", "REQ-COURSE-006", "Bắt buộc tên khóa học và mô tả", "P0", "Validation", ADMIN_IN, "CourseName rỗng; DescriptionHtml rỗng", "1. Submit với tên rỗng\n2. Submit với mô tả rỗng", "Hiển thị lỗi đúng trường; không tạo bản ghi.", "High")
add_case("Course", "REQ-COURSE-006", "Giới hạn tên khóa học 100 ký tự", "P1", "Boundary", ADMIN_IN, "Tên 100 ký tự và 101 ký tự", "1. Tạo với 100 ký tự\n2. Tạo với 101 ký tự", "100 ký tự được chấp nhận; 101 ký tự bị validation và không lưu.", "Medium")
add_case("Course", "REQ-COURSE-006", "Biên học phí khóa học", "P0", "Boundary", ADMIN_IN, "-1; 0; 100000000; 100000001", "Tạo/cập nhật lần lượt với từng giá trị học phí", "0 và 100.000.000 hợp lệ; -1 và 100.000.001 bị từ chối.", "High")
add_case("Course", "REQ-COURSE-006", "Biên tổng số buổi học", "P0", "Boundary", ADMIN_IN, "0; 1; 200; 201", "Tạo/cập nhật lần lượt với từng tổng số buổi", "1 và 200 hợp lệ; 0 và 201 bị từ chối.", "High")
add_case("Course", "REQ-COURSE-007", "Cập nhật khóa học và xử lý ảnh", "P1", "Regression", "Có khóa học QA với ảnh riêng; đã đăng nhập Admin.", "Đổi tên không chọn ảnh, sau đó chọn ảnh mới", "Lần 1 giữ nguyên URL ảnh cũ; lần 2 dùng ảnh mới và xóa file ảnh cũ không phải ảnh mặc định.", "High")
add_case("Course", "REQ-COURSE-008", "Xóa khóa học chưa có lớp", "P0", "Functional", "Có khóa học QA chưa liên kết lớp.", "Course QA độc lập", "1. Mở Delete\n2. Xác nhận xóa", "Xóa thành công, thông báo phù hợp và bản ghi không còn trong danh sách.", "Critical")
add_case("Course", "REQ-COURSE-008", "Chặn xóa khóa học đã có lớp", "P0", "Negative", DATA_READY, "Khóa IELTS có lớp IELTS-K02", "1. Mở Delete khóa IELTS\n2. Xác nhận", "Không xóa; hiển thị thông báo khóa học đã có lớp liên kết; dữ liệu lớp còn nguyên.", "Critical")

# Class: 14 cases
add_case("Class", "REQ-CLASS-001", "Hiển thị danh sách lớp và số ghi danh", "P1", "Smoke", ADMIN_IN, "Dữ liệu seed", "1. Mở /Class\n2. Đối chiếu lớp, khóa học và sĩ số hiện tại", "Danh sách đúng dữ liệu; lớp mới hơn trước; số ghi danh đúng.", "High")
add_case("Class", "REQ-CLASS-001", "Tìm lớp theo mã, tên và giảng viên", "P1", "Functional", ADMIN_IN, "GT-K01; Giao Tiếp; John", "Tìm lần lượt bằng ba từ khóa", "Mỗi từ khóa trả đúng lớp phù hợp với ClassCode, ClassName hoặc LecturerName.", "Medium")
add_case("Class", "REQ-CLASS-002", "Lọc lớp theo khóa học", "P1", "Functional", ADMIN_IN, "Course IELTS", "1. Chọn khóa IELTS\n2. Lọc", "Chỉ hiển thị lớp thuộc khóa IELTS.", "Medium")
add_case("Class", "REQ-CLASS-002", "Kết hợp tìm kiếm và lọc khóa học", "P2", "Functional", ADMIN_IN, "search=K02; course=IELTS", "Áp dụng đồng thời từ khóa và khóa học", "Kết quả thỏa cả hai điều kiện; filter vẫn hiển thị giá trị đang chọn.", "Medium")
add_case("Class", "REQ-CLASS-002", "Phân trang lớp 5 bản ghi và giữ bộ lọc", "P2", "Regression", "Có ít nhất 7 lớp cùng khóa; đã lọc khóa học.", "page=2", "1. Chuyển sang trang 2\n2. Quay lại trang 1", "Mỗi trang tối đa 5 bản ghi và tham số search/courseId được giữ.", "Medium")
add_case("Class", "REQ-CLASS-003", "Xem chi tiết lớp và học viên ghi danh", "P1", "Functional", DATA_READY, "GT-K01", "1. Mở Details GT-K01\n2. Kiểm tra khóa học và học viên", "Hiển thị đúng khóa học, thông tin lớp và các học viên seed đã ghi danh.", "High")
add_case("Class", "REQ-CLASS-004", "Tạo lớp hợp lệ", "P0", "Smoke", ADMIN_IN, "QA-CLASS-01; lớp QA; khóa active; lịch học; start date; capacity=20", "1. Mở Create\n2. Nhập dữ liệu\n3. Lưu", "Tạo thành công, thông báo phù hợp và lớp xuất hiện trong danh sách.", "Critical")
add_case("Class", "REQ-CLASS-005", "Chặn ClassCode trùng", "P0", "Negative", DATA_READY, "ClassCode=GT-K01", "1. Mở Create\n2. Nhập mã GT-K01 và dữ liệu hợp lệ\n3. Lưu", "Không tạo; hiển thị 'Mã lớp này đã tồn tại trên hệ thống'.", "Critical")
add_case("Class", "REQ-CLASS-006", "Kiểm tra trường lớp bắt buộc", "P0", "Validation", ADMIN_IN, "Bỏ trống lần lượt ClassCode, ClassName, CourseId, Schedule, StartDate", "Submit form cho từng trường hợp thiếu", "Hiển thị validation tương ứng và không tạo lớp.", "High")
add_case("Class", "REQ-CLASS-006", "Giới hạn độ dài dữ liệu lớp", "P1", "Boundary", ADMIN_IN, "ClassCode 20/21; ClassName 100/101; Lecturer 100/101; Schedule 100/101; Room 50/51", "Thử lần lượt giá trị tại biên và vượt biên", "Giá trị tại biên được chấp nhận; vượt biên bị từ chối đúng trường.", "Medium")
add_case("Class", "REQ-CLASS-006", "Biên sĩ số tối đa", "P0", "Boundary", ADMIN_IN, "0; 1; 100; 101", "Tạo/cập nhật lớp lần lượt với từng MaxCapacity", "1 và 100 hợp lệ; 0 và 101 bị validation.", "High")
add_case("Class", "REQ-CLASS-004", "Dropdown tạo lớp chỉ có khóa học active", "P1", "Functional", "Có một Course IsActive=false và một Course active.", "Hai khóa học test", "1. Mở Class/Create\n2. Kiểm tra danh sách khóa học", "Chỉ khóa active xuất hiện; khóa inactive không thể chọn từ UI.", "High")
add_case("Class", "REQ-CLASS-007", "Cập nhật thông tin lớp", "P1", "Functional", "Có lớp QA chưa có xung đột.", "Đổi tên, giảng viên, lịch, phòng, ngày và capacity", "1. Mở Edit\n2. Thay đổi dữ liệu\n3. Lưu\n4. Mở Details", "Cập nhật thành công và toàn bộ giá trị mới được hiển thị đúng.", "High")
add_case("Class", "REQ-CLASS-008", "Xóa lớp chưa có ghi danh và chặn lớp đã có ghi danh", "P0", "Regression", "Có lớp QA rỗng và lớp GT-K01 có ghi danh.", "Hai lớp test", "1. Xóa lớp QA rỗng\n2. Thử xóa GT-K01", "Lớp rỗng bị xóa; GT-K01 không bị xóa và có thông báo đã có học viên đăng ký.", "Critical")

# Enrollment: 18 cases
add_case("Enrollment", "REQ-ENR-001", "Hiển thị danh sách ghi danh mới nhất", "P1", "Smoke", ADMIN_IN, "Dữ liệu seed", "1. Mở /Enrollment\n2. Đối chiếu học viên, lớp, học phí và trạng thái", "Danh sách đúng dữ liệu và EnrollmentId lớn hơn nằm trước.", "High")
add_case("Enrollment", "REQ-ENR-001", "Tìm ghi danh theo học viên hoặc lớp", "P1", "Functional", ADMIN_IN, "Nguyễn Văn An; SV2026001; GT-K01; Giao Tiếp K01", "Tìm lần lượt bằng tên, mã học viên, mã lớp và tên lớp", "Mỗi lần trả các ghi danh chứa từ khóa ở trường tương ứng.", "Medium")
add_case("Enrollment", "REQ-ENR-001", "Lọc ghi danh theo lớp", "P1", "Functional", ADMIN_IN, "classId của GT-K01", "Chọn lớp GT-K01 và lọc", "Chỉ hiển thị ghi danh thuộc GT-K01.", "Medium")
add_case("Enrollment", "REQ-ENR-001", "Lọc ghi danh theo trạng thái học phí", "P1", "Functional", ADMIN_IN, "DaNopDu và ConNo", "Lọc lần lượt theo hai trạng thái", "Mỗi lần chỉ hiển thị bản ghi có PaymentStatus tương ứng.", "Medium")
add_case("Enrollment", "REQ-ENR-001", "Kết hợp search, lớp và trạng thái học phí", "P2", "Functional", ADMIN_IN, "SV2026001 + GT-K01 + DaNopDu", "Áp dụng đồng thời ba điều kiện", "Kết quả thỏa tất cả điều kiện và filter giữ nguyên giá trị.", "Medium")
add_case("Enrollment", "REQ-ENR-001", "Phân trang ghi danh 5 bản ghi", "P2", "Regression", "Có ít nhất 7 ghi danh phù hợp bộ lọc.", "page=1 và page=2", "Chuyển trang khi đang áp dụng bộ lọc", "Tối đa 5 bản ghi/trang; không trùng/mất và tham số lọc được giữ.", "Medium")
add_case("Enrollment", "REQ-ENR-002", "Xem chi tiết ghi danh cùng dữ liệu liên kết", "P1", "Functional", DATA_READY, "Ghi danh Nguyễn Văn An - GT-K01", "Mở Details của ghi danh", "Hiển thị đúng student, class/course, học phí, các bản ghi điểm danh và điểm số.", "High")
add_case("Enrollment", "REQ-ENR-003", "Tạo ghi danh hợp lệ với giá trị mặc định", "P0", "Smoke", "Có học viên và lớp còn chỗ, chưa ghi danh cặp này.", "Student QA; Class QA", "1. Mở Create\n2. Chọn học viên/lớp\n3. Lưu", "Tạo thành công; mặc định PaidAmount=0, PaymentStatus=ConNo, LearningStatus=DangHoc.", "Critical")
add_case("Enrollment", "REQ-ENR-004", "Tự điền học phí chuẩn khi chọn lớp", "P1", "Functional", ADMIN_IN, "GT-K01 có BaseTuitionFee=3500000", "1. Mở Create\n2. Chọn GT-K01\n3. Quan sát ActualFee", "API trả success=true, fee=3500000 và form điền đúng học phí.", "High")
add_case("Enrollment", "REQ-ENR-005", "Chặn ghi danh trùng học viên trong cùng lớp", "P0", "Negative", DATA_READY, "SV2026001 + GT-K01", "Tạo lại cặp học viên/lớp đã có", "Không tạo; hiển thị thông báo học viên đã được ghi danh vào lớp.", "Critical")
add_case("Enrollment", "REQ-ENR-006", "Chặn ghi danh khi lớp đủ sĩ số", "P0", "Boundary", "Lớp QA có MaxCapacity bằng đúng số ghi danh hiện tại.", "Thêm một học viên mới", "Submit ghi danh mới", "Không tạo; hiển thị lớp đã đủ sĩ số tối đa.", "Critical")
add_case("Enrollment", "REQ-ENR-007", "Tự chuyển DaNopDu khi đã nộp đủ hoặc dư", "P0", "Business Rule", "Cặp student/class chưa tồn tại và lớp còn chỗ.", "ActualFee=3500000; PaidAmount=3500000 rồi 4000000", "Tạo hai ghi danh test độc lập", "Cả hai bản ghi có PaymentStatus=DaNopDu.", "Critical")
add_case("Enrollment", "REQ-ENR-007", "Tự chuyển ConNo khi nộp thiếu", "P0", "Business Rule", "Cặp student/class chưa tồn tại và lớp còn chỗ.", "ActualFee=3500000; PaidAmount=3499999", "Tạo ghi danh", "Bản ghi có PaymentStatus=ConNo, bất kể giá trị PaymentStatus gửi từ client.", "Critical")
add_case("Enrollment", "REQ-ENR-007", "ActualFee bằng 0 vẫn là ConNo", "P1", "Boundary", "Cặp student/class chưa tồn tại.", "ActualFee=0; PaidAmount=0 hoặc 1000", "Tạo ghi danh cho từng bộ dữ liệu", "PaymentStatus=ConNo vì ActualFee không lớn hơn 0.", "High")
add_case("Enrollment", "REQ-ENR-008", "Biên học phí thực tế và số tiền đã nộp", "P0", "Boundary", "Có dữ liệu tạo ghi danh hợp lệ khác.", "-1; 0; 100000000; 100000001 cho mỗi trường", "Submit lần lượt các giá trị tại/vượt biên", "0 và 100.000.000 hợp lệ; số âm và trên 100.000.000 bị từ chối.", "High")
add_case("Enrollment", "REQ-ENR-009", "Cập nhật ghi danh, trạng thái học và học phí", "P1", "Regression", "Có ghi danh QA chưa có attendance/grade.", "LearningStatus=BaoLuu; tăng PaidAmount đến đủ", "1. Mở Edit\n2. Đổi trạng thái và tiền đã nộp\n3. Lưu", "Thông tin cập nhật; LearningStatus=BaoLuu và PaymentStatus tự đổi thành DaNopDu.", "High")
add_case("Enrollment", "REQ-ENR-010", "Xóa ghi danh chưa có điểm danh/điểm", "P0", "Functional", "Có ghi danh QA không có Attendance/Grade.", "Enrollment QA", "1. Mở Delete\n2. Xác nhận", "Xóa thành công; sĩ số lớp giảm một và bản ghi biến mất.", "Critical")
add_case("Enrollment", "REQ-ENR-010", "Chặn xóa ghi danh có điểm danh hoặc điểm", "P0", "Negative", DATA_READY, "Enrollment có Attendance; Enrollment có Grade", "Thử xóa từng ghi danh có dữ liệu liên kết", "Không xóa cả hai; hiển thị thông báo đã có điểm danh hoặc bảng điểm; dữ liệu con còn nguyên.", "Critical")

# Attendance: 12 cases
add_case("Attendance", "REQ-ATT-001", "Hiển thị danh sách điểm danh theo ngày và buổi mới nhất", "P1", "Smoke", ADMIN_IN, "Dữ liệu seed", "Mở /Attendance và đối chiếu thứ tự", "Danh sách đúng dữ liệu; ngày mới hơn trước, cùng ngày thì buổi lớn hơn trước.", "High")
add_case("Attendance", "REQ-ATT-001", "Tìm điểm danh theo học viên hoặc mã lớp", "P1", "Functional", ADMIN_IN, "Nguyễn Văn An; SV2026001; GT-K01", "Tìm lần lượt theo ba từ khóa", "Kết quả phù hợp FullName, StudentCode hoặc ClassCode.", "Medium")
add_case("Attendance", "REQ-ATT-001", "Lọc điểm danh theo lớp, buổi và ngày", "P1", "Functional", ADMIN_IN, "GT-K01; session=1; ngày seed", "Áp dụng từng filter và sau đó kết hợp cả ba", "Kết quả đúng từng điều kiện và giao của các điều kiện khi kết hợp.", "Medium")
add_case("Attendance", "REQ-ATT-001", "Phân trang điểm danh 10 bản ghi và giữ bộ lọc", "P2", "Regression", "Có trên 10 bản ghi phù hợp.", "page=2", "Chuyển trang khi có search/filter", "Tối đa 10 bản ghi/trang; tham số search, classId, sessionNumber, date được giữ.", "Medium")
add_case("Attendance", "REQ-ATT-002", "Sổ điểm danh chỉ hiển thị học viên DangHoc", "P0", "Business Rule", "Một lớp có ghi danh DangHoc và BaoLuu/DaNghi.", "Chọn lớp test", "Mở TakeAttendance của lớp", "Chỉ học viên DangHoc xuất hiện; các trạng thái khác không xuất hiện.", "Critical")
add_case("Attendance", "REQ-ATT-003", "Mặc định học viên mới là có mặt", "P1", "Functional", "Lớp có học viên DangHoc chưa điểm danh buổi được chọn.", "session mới", "Mở TakeAttendance", "Mỗi học viên chưa có dữ liệu được tick IsPresent=true và Note rỗng.", "High")
add_case("Attendance", "REQ-ATT-003", "Nạp dữ liệu điểm danh đã tồn tại", "P1", "Regression", DATA_READY, "GT-K01; session=1", "Mở TakeAttendance cho lớp và buổi đã điểm danh", "Hiển thị đúng IsPresent, Note, AttendanceDate và giữ AttendanceId để cập nhật.", "High")
add_case("Attendance", "REQ-ATT-003", "Lưu điểm danh mới cho cả lớp", "P0", "Smoke", "Lớp có học viên DangHoc và buổi chưa được lưu.", "session=3; hôm nay; một có mặt, một vắng", "1. Chọn lớp/buổi/ngày\n2. Chỉnh trạng thái và note\n3. Lưu", "Mỗi học viên có một bản ghi mới đúng dữ liệu; thông báo lưu thành công.", "Critical")
add_case("Attendance", "REQ-ATT-003", "Cập nhật điểm danh đã có không tạo bản ghi trùng", "P0", "Regression", "Buổi điểm danh đã tồn tại.", "Đổi vắng thành có mặt và sửa note", "1. Mở lại sổ điểm danh\n2. Sửa dữ liệu\n3. Lưu", "Bản ghi cũ được cập nhật qua AttendanceId; tổng số bản ghi của học viên/buổi không tăng.", "Critical")
add_case("Attendance", "REQ-ATT-004", "Validation lớp, ngày và biên số buổi", "P0", "Boundary", ADMIN_IN, "ClassId=0; date rỗng; session=0,1,200,201", "Submit lần lượt các bộ dữ liệu", "Class/date thiếu bị từ chối; session 1 và 200 hợp lệ, 0 và 201 không hợp lệ.", "High")
add_case("Attendance", "REQ-ATT-005", "Giới hạn ghi chú điểm danh 255 ký tự", "P1", "Boundary", "Có sổ điểm danh hợp lệ.", "Note 255 và 256 ký tự", "Lưu lần lượt hai độ dài", "255 ký tự được lưu; 256 ký tự bị từ chối hoặc không được ghi vào database.", "Medium")
add_case("Attendance", "REQ-ATT-006", "Xóa bản ghi điểm danh", "P1", "Functional", "Có Attendance QA có thể xóa.", "Attendance QA", "1. Mở Delete\n2. Xác nhận", "Xóa thành công, thông báo phù hợp và bản ghi không còn trong danh sách.", "High")

# Grade: 15 cases
add_case("Grade", "REQ-GRADE-001", "Hiển thị danh sách điểm mới nhất", "P1", "Smoke", ADMIN_IN, "Dữ liệu seed", "Mở /Grade và đối chiếu dữ liệu", "Danh sách đúng student/class/course/exam/điểm; GradeId lớn hơn trước.", "High")
add_case("Grade", "REQ-GRADE-001", "Tìm điểm theo học viên hoặc mã lớp", "P1", "Functional", ADMIN_IN, "Trần Thị Mai; SV2026002; IELTS-K02", "Tìm lần lượt theo ba từ khóa", "Kết quả phù hợp FullName, StudentCode hoặc ClassCode.", "Medium")
add_case("Grade", "REQ-GRADE-001", "Lọc điểm theo lớp và kỳ kiểm tra", "P1", "Functional", ADMIN_IN, "IELTS-K02; PlacementTest", "Lọc riêng từng điều kiện rồi kết hợp", "Kết quả đúng điều kiện và giữ lựa chọn filter.", "Medium")
add_case("Grade", "REQ-GRADE-001", "Phân trang điểm 10 bản ghi và giữ bộ lọc", "P2", "Regression", "Có trên 10 bản ghi phù hợp.", "page=2", "Chuyển trang khi đang search/filter", "Tối đa 10 bản ghi/trang, không trùng/mất và giữ tham số.", "Medium")
add_case("Grade", "REQ-GRADE-002", "Sổ điểm chỉ hiển thị học viên DangHoc", "P0", "Business Rule", "Một lớp có ghi danh DangHoc và BaoLuu/DaNghi.", "Chọn lớp test", "Mở EnterGrades", "Chỉ học viên DangHoc xuất hiện.", "Critical")
add_case("Grade", "REQ-GRADE-002", "Nạp kỳ mặc định và điểm đã tồn tại", "P1", "Regression", DATA_READY, "Không truyền examType; sau đó GT-K01/GiuaKy", "1. Mở EnterGrades không có examType\n2. Mở lớp/kỳ đã có điểm", "Mặc định GiuaKy; lần 2 hiển thị đúng các kỹ năng, Overall, feedback và GradeId.", "High")
add_case("Grade", "REQ-GRADE-003", "Tạo bản ghi điểm mới khi có dữ liệu", "P0", "Smoke", "Học viên DangHoc chưa có điểm ở kỳ test.", "Một kỹ năng hoặc feedback", "Nhập dữ liệu và lưu", "Tạo đúng một Grade, thông báo thành công và mở lại thấy dữ liệu.", "Critical")
add_case("Grade", "REQ-GRADE-003", "Không tạo Grade cho dòng hoàn toàn trống", "P0", "Negative", "Học viên DangHoc chưa có điểm ở kỳ test.", "Bốn điểm null; feedback rỗng", "Để trống dòng học viên và lưu sổ điểm", "Không tạo bản ghi Grade cho học viên đó.", "High")
add_case("Grade", "REQ-GRADE-003", "Cập nhật Grade đã có không tạo trùng", "P0", "Regression", "Đã có GradeId cho học viên/kỳ.", "Sửa ReadingScore và feedback", "1. Mở sổ điểm\n2. Sửa\n3. Lưu", "Grade cũ được cập nhật; không tăng số bản ghi cho học viên/kỳ.", "Critical")
add_case("Grade", "REQ-GRADE-004", "Biên điểm kỹ năng 0-990", "P0", "Boundary", "Có sổ điểm hợp lệ.", "-0.1; 0; 990; 990.1 cho từng kỹ năng", "Submit lần lượt giá trị biên và vượt biên", "0 và 990 hợp lệ; giá trị nhỏ hơn 0 hoặc lớn hơn 990 bị từ chối.", "High")
add_case("Grade", "REQ-GRADE-005", "Validation kỳ kiểm tra và nhận xét", "P1", "Boundary", "Có sổ điểm hợp lệ.", "ExamType rỗng/50/51 ký tự; feedback 500/501 ký tự", "Submit lần lượt các dữ liệu", "Rỗng/vượt giới hạn bị từ chối; giá trị đúng giới hạn được lưu.", "Medium")
add_case("Grade", "REQ-GRADE-006", "Tính tổng TOEIC từ Listening và Reading", "P0", "Business Rule", "CourseName hoặc ExamType chứa TOEIC.", "L=400,R=350; L=400,R=null; cả hai null", "Lưu từng bộ dữ liệu", "Overall lần lượt 750, 400 và null; Writing/Speaking không ảnh hưởng tổng TOEIC.", "Critical")
add_case("Grade", "REQ-GRADE-007", "Làm tròn IELTS tại các ngưỡng band", "P0", "Business Rule", "CourseName hoặc ExamType chứa IELTS.", "Average=6.24, 6.25, 6.74, 6.75", "Nhập bộ điểm tạo từng average rồi lưu", "Overall lần lượt 6.0, 6.5, 6.5, 7.0.", "Critical")
add_case("Grade", "REQ-GRADE-008", "Tính trung bình khóa học thường", "P1", "Business Rule", "Course/ExamType không chứa IELTS hoặc TOEIC.", "L=7.2,R=6.4,W=null,S=8.1", "Lưu điểm", "Overall là trung bình ba kỹ năng đã nhập, làm tròn một chữ số: 7.2.", "High")
add_case("Grade", "REQ-GRADE-009", "Xóa bản ghi điểm", "P1", "Functional", "Có Grade QA có thể xóa.", "Grade QA", "1. Mở Delete\n2. Xác nhận", "Xóa thành công, thông báo phù hợp và bản ghi không còn trong danh sách/sổ điểm.", "High")

# Common: 6 cases
add_case("Common", "REQ-COMMON-001", "Chưa đăng nhập bị chuyển đến Login từ mọi module quản trị", "P0", "Authorization", "Xóa cookie/phiên đăng nhập.", "/Course, /Class, /Enrollment, /Attendance, /Grade", "Mở trực tiếp từng URL", "Mỗi request chuyển đến /Account/Login và giữ ReturnUrl nội bộ.", "Critical")
add_case("Common", "REQ-COMMON-002", "POST thay đổi dữ liệu thiếu anti-forgery token bị từ chối", "P0", "Security", "Có phiên Admin hợp lệ.", "POST Create/Edit/Delete/Save không có token", "Gửi request trực tiếp không kèm token tới từng nhóm action", "Request bị từ chối (HTTP 400); database không thay đổi.", "Critical")
add_case("Common", "REQ-COMMON-003", "Xử lý id rỗng hoặc không tồn tại", "P1", "Negative", ADMIN_IN, "id rỗng; id=999999 cho Details/Edit/Delete", "Thử trên Course, Class, Enrollment, Attendance, Grade", "Trả Not Found phù hợp; không lộ exception stack trace.", "High")
add_case("Common", "REQ-COMMON-004", "Build solution thành công", "P0", "Build", "Đã restore package; có .NET 10 SDK.", "dotnet build Assignment.slnx --no-restore", "Chạy lệnh build tại repo root", "Exit code 0 và không có compile error.", "Critical")
add_case("Common", "REQ-COMMON-004", "Khởi tạo database và seed account/data", "P0", "Integration", "SQL Server khả dụng; database test chưa tồn tại; connection string hợp lệ.", "EnglishCenterDb test sạch", "1. Chạy app\n2. Kiểm tra schema/data\n3. Đăng nhập bốn tài khoản seed", "Database được tạo một lần; role/user/student/course/class/enrollment/attendance/grade seed có mặt và tài khoản đăng nhập được.", "Critical")
add_case("Common", "REQ-COMMON-005", "Kiểm tra trình duyệt, bố cục và thao tác bàn phím cơ bản", "P2", "UI", "Ứng dụng chạy với dữ liệu seed.", "Chrome/Edge; 1366x768 và 1920x1080", "Đi qua luồng chính bằng chuột và bàn phím; resize hai viewport", "Không có nội dung chồng/tràn; focus nhìn thấy; label/validation đọc được; thao tác chính hoàn tất trên cả hai trình duyệt.", "Medium")


assert len(CASES) == 90, f"Expected 90 test cases, got {len(CASES)}"


def set_title(ws, title: str, subtitle: str, width: int):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    ws.cell(1, 1, title)
    ws.cell(1, 1).font = Font(size=16, bold=True, color=WHITE)
    ws.cell(1, 1).fill = PatternFill("solid", fgColor=NAVY)
    ws.cell(1, 1).alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    ws.cell(2, 1, subtitle)
    ws.cell(2, 1).font = Font(italic=True, color="595959")
    ws.cell(2, 1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 30


def style_header(ws, row: int, start_col: int, end_col: int):
    for cell in ws.iter_cols(min_col=start_col, max_col=end_col, min_row=row, max_row=row):
        target = cell[0]
        target.font = Font(bold=True, color=WHITE)
        target.fill = PatternFill("solid", fgColor=BLUE)
        target.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        target.border = Border(top=THIN_GRAY, bottom=THIN_GRAY, left=THIN_GRAY, right=THIN_GRAY)
    ws.row_dimensions[row].height = 32


def style_data(ws, min_row: int, max_row: int, min_col: int, max_col: int):
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, min_col=min_col, max_col=max_col):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(top=THIN_GRAY, bottom=THIN_GRAY, left=THIN_GRAY, right=THIN_GRAY)


def add_table(ws, ref: str, name: str):
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def add_list_validation(ws, cell_range: str, values: list[str]):
    formula = '"' + ",".join(values) + '"'
    validation = DataValidation(type="list", formula1=formula, allow_blank=True)
    validation.error = "Vui lòng chọn một giá trị trong danh sách."
    validation.errorTitle = "Giá trị không hợp lệ"
    validation.prompt = "Chọn giá trị chuẩn."
    validation.promptTitle = "Danh sách lựa chọn"
    ws.add_data_validation(validation)
    validation.add(cell_range)


def build_test_plan(wb: Workbook):
    ws = wb.active
    ws.title = "Test Plan"
    set_title(ws, "TEST PLAN - ENGLISH CENTER MVC", "Baseline QA theo Waterfall | Cập nhật: " + date.today().isoformat(), 4)
    headers = ["Nhóm", "Hạng mục", "Nội dung", "Owner/Trạng thái"]
    ws.append([])
    ws.append(headers)
    rows = [
        ["Thông tin", "Dự án", "Hệ thống quản lý học viên trung tâm tiếng Anh", "Project Lead"],
        ["Thông tin", "Phiên bản", "QA Baseline v1.0", "Approved"],
        ["Phạm vi", "In scope", "Auth; Course; Class; Enrollment; Attendance; Grade; Search/Filter/Pagination; Authorization", "QA"],
        ["Phạm vi", "Out of scope", "Online payment; email/SMS; public API; performance/load; production deployment", "Agreed"],
        ["Môi trường", "Stack", ".NET 10; ASP.NET Core MVC; EF Core; SQL Server; cookie auth", "Dev"],
        ["Môi trường", "Trình duyệt", "Chrome và Edge stable; desktop 1366x768 trở lên", "QA"],
        ["Môi trường", "Build", "dotnet build Assignment.slnx --no-restore", "Exit code 0"],
        ["Tài khoản", "Admin", "admin / admin@123", "Seed"],
        ["Tài khoản", "GiaoVu", "giaovu01 / giaovu@123", "Seed"],
        ["Tài khoản", "SinhVien 1", "sv2026001 / 123456", "Seed"],
        ["Tài khoản", "SinhVien 2", "sv2026002 / 123456", "Seed"],
        ["Entry", "Requirement", "Requirement baseline được duyệt và mỗi requirement có REQ ID.", "Required"],
        ["Entry", "Build & data", "Build pass; database test/seed sẵn sàng; không có blocker môi trường.", "Required"],
        ["Exit", "Smoke", "100% smoke test Pass.", "Required"],
        ["Exit", "Functional", ">= 95% test đã chạy Pass; N/A không tính vào mẫu số.", "Required"],
        ["Exit", "Defect", "Không còn bug Blocker/Critical mở; bug còn lại có owner/quyết định.", "Required"],
        ["Quy trình", "Thứ tự", "Build -> Smoke -> Functional -> Authorization -> Regression -> Retest -> Sign-off", "QA Lead"],
        ["Quy trình", "Evidence", "Cập nhật Actual Result, Status, Tester, Date; mỗi Fail phải có Bug Log liên kết.", "QA"],
        ["Timeline", "Tuần 1", "Requirements + Analysis", "Gate 1"],
        ["Timeline", "Tuần 2", "Design + Course/Class implementation", "Gate 2"],
        ["Timeline", "Tuần 3", "Enrollment/Attendance/Grade/Auth polish + module smoke", "Gate 3"],
        ["Timeline", "Tuần 4", "Full test, regression, handover và sign-off", "Gate 4"],
        ["Rủi ro", "Database", "EnsureCreated chưa hỗ trợ versioned migration; reset dữ liệu test có kiểm soát.", "Medium"],
        ["Rủi ro", "Automation", "Chưa có test project; ưu tiên automation cho công thức điểm và authorization.", "High"],
        ["Rủi ro", "Upload", "Cần xác nhận giới hạn MIME/extension/size phía server trước production.", "High"],
    ]
    for row in rows:
        ws.append(row)
    style_header(ws, 4, 1, 4)
    style_data(ws, 5, ws.max_row, 1, 4)
    add_table(ws, f"A4:D{ws.max_row}", "TestPlanTable")
    ws.freeze_panes = "A5"
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 92
    ws.column_dimensions["D"].width = 22
    ws.sheet_view.showGridLines = False


def build_requirements(wb: Workbook):
    ws = wb.create_sheet("Requirements")
    set_title(ws, "REQUIREMENT BASELINE", "Yêu cầu chức năng và phi chức năng dùng để truy vết test case.", 6)
    ws.append([])
    headers = ["REQ ID", "Module", "Mô tả yêu cầu", "Ưu tiên", "Nguồn", "Trạng thái"]
    ws.append(headers)
    for row in REQUIREMENTS:
        ws.append(row)
    style_header(ws, 4, 1, 6)
    style_data(ws, 5, ws.max_row, 1, 6)
    add_table(ws, f"A4:F{ws.max_row}", "RequirementsTable")
    add_list_validation(ws, f"D5:D{ws.max_row}", ["P0", "P1", "P2", "P3"])
    add_list_validation(ws, f"F5:F{ws.max_row}", ["Draft", "Baseline v1.0", "Changed", "Deprecated"])
    ws.freeze_panes = "A5"
    widths = [18, 16, 78, 12, 34, 18]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + idx)].width = width
    ws.sheet_view.showGridLines = False


def build_test_cases(wb: Workbook):
    ws = wb.create_sheet("Test Cases")
    set_title(ws, "TEST CASES", f"Bộ test baseline gồm {len(CASES)} case. Các cột kết quả để QA cập nhật khi chạy test.", 16)
    ws.append([])
    headers = [
        "TC ID", "Module", "REQ ID", "Tiêu đề", "Ưu tiên", "Loại", "Tiền điều kiện",
        "Dữ liệu kiểm thử", "Các bước thực hiện", "Kết quả mong đợi", "Kết quả thực tế",
        "Trạng thái", "Mức độ nếu lỗi", "Người kiểm thử", "Ngày chạy", "Ghi chú",
    ]
    ws.append(headers)
    for case in CASES:
        ws.append(list(case.values()))
    style_header(ws, 4, 1, 16)
    style_data(ws, 5, ws.max_row, 1, 16)
    add_table(ws, f"A4:P{ws.max_row}", "TestCasesTable")
    add_list_validation(ws, f"E5:E{ws.max_row}", ["P0", "P1", "P2", "P3"])
    add_list_validation(ws, f"F5:F{ws.max_row}", ["Smoke", "Functional", "Negative", "Validation", "Boundary", "Business Rule", "Authorization", "Security", "Regression", "Build", "Integration", "UI"])
    add_list_validation(ws, f"L5:L{ws.max_row}", ["Chưa chạy", "Pass", "Fail", "Blocked", "N/A"])
    add_list_validation(ws, f"M5:M{ws.max_row}", ["Blocker", "Critical", "High", "Medium", "Low"])
    ws.conditional_formatting.add(f"L5:L{ws.max_row}", FormulaRule(formula=["$L5=\"Pass\""], fill=PatternFill("solid", fgColor=GREEN)))
    ws.conditional_formatting.add(f"L5:L{ws.max_row}", FormulaRule(formula=["$L5=\"Fail\""], fill=PatternFill("solid", fgColor=RED)))
    ws.conditional_formatting.add(f"L5:L{ws.max_row}", FormulaRule(formula=["$L5=\"Blocked\""], fill=PatternFill("solid", fgColor=ORANGE)))
    ws.conditional_formatting.add(f"L5:L{ws.max_row}", FormulaRule(formula=["$L5=\"Chưa chạy\""], fill=PatternFill("solid", fgColor=YELLOW)))
    ws.freeze_panes = "A5"
    widths = [16, 14, 18, 42, 11, 16, 38, 36, 54, 54, 44, 14, 17, 18, 14, 34]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + idx) if idx <= 26 else "A"].width = width
    for row in range(5, ws.max_row + 1):
        ws.row_dimensions[row].height = 72
    ws.sheet_view.showGridLines = False


def build_rtm(wb: Workbook):
    ws = wb.create_sheet("RTM")
    set_title(ws, "REQUIREMENTS TRACEABILITY MATRIX", "Liên kết Requirement -> Test Case -> kết quả chạy test.", 10)
    ws.append([])
    headers = ["REQ ID", "Module", "Yêu cầu", "Test Case IDs", "Số TC", "Pass", "Fail", "Blocked", "Coverage", "Kết quả mới nhất"]
    ws.append(headers)
    by_req: dict[str, list[str]] = defaultdict(list)
    for case in CASES:
        by_req[case["req"]].append(case["id"])
    for index, req in enumerate(REQUIREMENTS, start=5):
        req_id, module, description = req[0], req[1], req[2]
        ids = ", ".join(by_req[req_id])
        ws.append([
            req_id,
            module,
            description,
            ids,
            f'=COUNTIF(\'Test Cases\'!$C:$C,A{index})',
            f'=COUNTIFS(\'Test Cases\'!$C:$C,A{index},\'Test Cases\'!$L:$L,"Pass")',
            f'=COUNTIFS(\'Test Cases\'!$C:$C,A{index},\'Test Cases\'!$L:$L,"Fail")',
            f'=COUNTIFS(\'Test Cases\'!$C:$C,A{index},\'Test Cases\'!$L:$L,"Blocked")',
            f'=IF(E{index}=0,"Chưa bao phủ","Đã bao phủ")',
            f'=IF(E{index}=0,"No TC",IF(G{index}>0,"Fail",IF(H{index}>0,"Blocked",IF(F{index}=E{index},"Pass","Chưa hoàn tất"))))',
        ])
    style_header(ws, 4, 1, 10)
    style_data(ws, 5, ws.max_row, 1, 10)
    add_table(ws, f"A4:J{ws.max_row}", "RTMTable")
    ws.conditional_formatting.add(f"J5:J{ws.max_row}", FormulaRule(formula=["$J5=\"Pass\""], fill=PatternFill("solid", fgColor=GREEN)))
    ws.conditional_formatting.add(f"J5:J{ws.max_row}", FormulaRule(formula=["$J5=\"Fail\""], fill=PatternFill("solid", fgColor=RED)))
    ws.conditional_formatting.add(f"J5:J{ws.max_row}", FormulaRule(formula=["$J5=\"Blocked\""], fill=PatternFill("solid", fgColor=ORANGE)))
    ws.freeze_panes = "A5"
    widths = [18, 14, 66, 58, 11, 11, 11, 12, 17, 18]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + idx)].width = width
    ws.sheet_view.showGridLines = False


def build_smoke(wb: Workbook):
    ws = wb.create_sheet("Smoke Checklist")
    set_title(ws, "SMOKE CHECKLIST", "Chạy sau mỗi build/deploy. Exit gate yêu cầu 100% Pass.", 7)
    ws.append([])
    headers = ["Smoke ID", "Khu vực", "Kiểm tra", "Kết quả mong đợi", "Trạng thái", "Người chạy", "Ghi chú"]
    ws.append(headers)
    smoke = [
        ["SMK-001", "Build", "Build solution", "Exit code 0, không có compile error"],
        ["SMK-002", "Startup", "Ứng dụng khởi động và kết nối database", "Trang Login mở được, không có lỗi startup"],
        ["SMK-003", "Seed", "Đăng nhập bốn tài khoản seed", "Cả bốn tài khoản xác thực đúng"],
        ["SMK-004", "Auth", "Admin đăng nhập và đăng xuất", "Tạo/hủy phiên đúng"],
        ["SMK-005", "Authorization", "SinhVien mở /Course", "Chuyển đến Denied"],
        ["SMK-006", "Course", "Tạo khóa học tối thiểu hợp lệ", "Tạo và hiển thị trong danh sách"],
        ["SMK-007", "Class", "Tạo lớp cho khóa học QA", "Tạo và hiển thị trong danh sách"],
        ["SMK-008", "Enrollment", "Ghi danh học viên vào lớp còn chỗ", "Tạo ghi danh và tính PaymentStatus đúng"],
        ["SMK-009", "Attendance", "Lưu điểm danh một buổi", "Tạo bản ghi và mở lại thấy dữ liệu"],
        ["SMK-010", "Grade", "Lưu điểm một kỳ", "Tạo bản ghi và Overall đúng loại khóa"],
        ["SMK-011", "Search", "Tìm kiếm trên danh sách chính", "Kết quả lọc đúng, không lỗi"],
        ["SMK-012", "Delete guard", "Thử xóa khóa/lớp/ghi danh có liên kết", "Bị chặn và dữ liệu còn nguyên"],
        ["SMK-013", "Navigation", "Đi qua menu Course -> Grade", "Các trang mở được, không có link hỏng"],
        ["SMK-014", "Error", "Mở Details với id không tồn tại", "Trả Not Found, không lộ stack trace"],
    ]
    for row in smoke:
        ws.append(row + ["Chưa chạy", "", ""])
    style_header(ws, 4, 1, 7)
    style_data(ws, 5, ws.max_row, 1, 7)
    add_table(ws, f"A4:G{ws.max_row}", "SmokeTable")
    add_list_validation(ws, f"E5:E{ws.max_row}", ["Chưa chạy", "Pass", "Fail", "Blocked", "N/A"])
    ws.freeze_panes = "A5"
    widths = [15, 18, 48, 56, 14, 18, 36]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + idx)].width = width
    ws.sheet_view.showGridLines = False


def build_regression(wb: Workbook):
    ws = wb.create_sheet("Regression Checklist")
    set_title(ws, "REGRESSION CHECKLIST", "Chọn checklist theo vùng thay đổi và chạy trước release/sign-off.", 8)
    ws.append([])
    headers = ["REG ID", "Vùng thay đổi", "Trigger", "Kiểm tra bắt buộc", "TC tham chiếu", "Trạng thái", "Người chạy", "Ghi chú"]
    ws.append(headers)
    rows = [
        ["REG-001", "Auth", "Đổi login/cookie/claims", "Login 3 role, sai mật khẩu, ReturnUrl, logout, denied", "TC-AUTH-001..011"],
        ["REG-002", "Authorization", "Đổi role/policy/controller", "Anonymous redirect; SinhVien denied; Admin/GiaoVu allowed", "TC-AUTH-011; TC-COMMON-001"],
        ["REG-003", "Course", "Đổi model/controller/view Course", "CRUD, search, page, validation, ảnh, delete guard", "TC-COURSE-001..014"],
        ["REG-004", "Upload", "Đổi upload/static files", "Ảnh mới, ảnh mặc định, giữ/thay ảnh, đường dẫn hiển thị", "TC-COURSE-006..012"],
        ["REG-005", "Class", "Đổi model/controller/view Class", "CRUD, duplicate code, active course, capacity, delete guard", "TC-CLASS-001..014"],
        ["REG-006", "Enrollment", "Đổi ghi danh/học phí", "Duplicate, capacity, fee API, payment status, range, delete guard", "TC-ENROLLMENT-001..018"],
        ["REG-007", "Attendance", "Đổi sổ điểm danh", "Active students, default present, create/update, validation, delete", "TC-ATTENDANCE-001..012"],
        ["REG-008", "Grade", "Đổi sổ điểm/công thức", "Create/update/blank row, range, IELTS, TOEIC, average, delete", "TC-GRADE-001..015"],
        ["REG-009", "Search/filter", "Đổi query hoặc PagedList", "Search, combined filters, no result, giữ query khi chuyển trang", "Các TC list/search/page"],
        ["REG-010", "Data model", "Đổi entity/DbContext", "Startup DB, seed, CRUD liên quan và delete constraints", "TC-COMMON-005 + module CRUD"],
        ["REG-011", "Validation", "Đổi DataAnnotations/view form", "Required, min/max, max length, invalid POST giữ dữ liệu", "Các TC Validation/Boundary"],
        ["REG-012", "Navigation/layout", "Đổi shared view/CSS/JS", "Menu, validation scripts, desktop viewport, thao tác bàn phím", "TC-COMMON-006"],
        ["REG-013", "Package/runtime", "Nâng .NET/NuGet", "Restore, build, startup, auth, full smoke", "TC-COMMON-004; SMK-001..014"],
        ["REG-014", "Bug fix", "Retest một bug", "Chạy linked TC, case biên lân cận và module smoke", "Theo Bug Log"],
    ]
    for row in rows:
        ws.append(row + ["Chưa chạy", "", ""])
    style_header(ws, 4, 1, 8)
    style_data(ws, 5, ws.max_row, 1, 8)
    add_table(ws, f"A4:H{ws.max_row}", "RegressionTable")
    add_list_validation(ws, f"F5:F{ws.max_row}", ["Chưa chạy", "Pass", "Fail", "Blocked", "N/A"])
    ws.freeze_panes = "A5"
    widths = [14, 20, 32, 62, 34, 14, 18, 34]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + idx)].width = width
    ws.sheet_view.showGridLines = False


def build_bug_log(wb: Workbook):
    ws = wb.create_sheet("Bug Log")
    set_title(ws, "BUG LOG", "Mỗi test Fail cần một bug liên kết. Có sẵn 30 dòng nhập liệu.", 13)
    ws.append([])
    headers = ["Bug ID", "Linked TC", "Module", "Severity", "Tiêu đề", "Bước tái hiện", "Kết quả mong đợi", "Kết quả thực tế", "Owner", "Trạng thái", "Ngày tạo", "Ngày retest", "Ghi chú"]
    ws.append(headers)
    for _ in range(30):
        ws.append([""] * len(headers))
    style_header(ws, 4, 1, 13)
    style_data(ws, 5, ws.max_row, 1, 13)
    add_table(ws, f"A4:M{ws.max_row}", "BugLogTable")
    add_list_validation(ws, f"C5:C{ws.max_row}", ["Auth", "Course", "Class", "Enrollment", "Attendance", "Grade", "Common"])
    add_list_validation(ws, f"D5:D{ws.max_row}", ["Blocker", "Critical", "High", "Medium", "Low"])
    add_list_validation(ws, f"J5:J{ws.max_row}", ["Open", "In Progress", "Ready to Retest", "Closed", "Rejected", "Deferred"])
    ws.freeze_panes = "A5"
    widths = [14, 18, 15, 14, 38, 52, 45, 45, 18, 18, 14, 14, 32]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + idx)].width = width
    ws.sheet_view.showGridLines = False


def build_summary(wb: Workbook):
    ws = wb.create_sheet("Test Summary")
    set_title(ws, "TEST SUMMARY & SIGN-OFF", "Công thức tự cập nhật từ Test Cases và Bug Log khi mở bằng Excel/LibreOffice.", 8)
    ws.append([])
    ws.append(["Chỉ số", "Giá trị", "Tiêu chí", "Đánh giá", "", "Module", "Tổng TC", "Pass/Fail/Blocked/Chưa chạy"])
    metrics = [
        ["Tổng test case", "=COUNTA('Test Cases'!A:A)-3", "90 baseline", "=IF(B5=90,\"OK\",\"Kiểm tra\")"],
        ["Pass", '=COUNTIF(\'Test Cases\'!L:L,"Pass")', "Theo kết quả chạy", ""],
        ["Fail", '=COUNTIF(\'Test Cases\'!L:L,"Fail")', "Mỗi Fail có Bug Log", '=IF(B7=0,"OK","Có lỗi")'],
        ["Blocked", '=COUNTIF(\'Test Cases\'!L:L,"Blocked")', "Phải có nguyên nhân", '=IF(B8=0,"OK","Cần xử lý")'],
        ["Chưa chạy", '=COUNTIF(\'Test Cases\'!L:L,"Chưa chạy")', "0 trước sign-off", '=IF(B9=0,"OK","Chưa hoàn tất")'],
        ["N/A", '=COUNTIF(\'Test Cases\'!L:L,"N/A")', "Có lý do trong Note", ""],
        ["Pass rate", '=IFERROR(B6/(B5-B9-B10),0)', ">= 95%", '=IF(B11>=95%,"Đạt","Chưa đạt")'],
        ["Smoke pass", '=COUNTIF(\'Smoke Checklist\'!E:E,"Pass")', "14/14", '=IF(B12=14,"Đạt","Chưa đạt")'],
        ["Bug Blocker/Critical mở", '=COUNTIFS(\'Bug Log\'!D:D,"Blocker",\'Bug Log\'!J:J,"<>Closed",\'Bug Log\'!A:A,"<>")+COUNTIFS(\'Bug Log\'!D:D,"Critical",\'Bug Log\'!J:J,"<>Closed",\'Bug Log\'!A:A,"<>")', "0", '=IF(B13=0,"Đạt","Chưa đạt")'],
        ["Build", "Pass", "Exit code 0", "Đạt"],
        ["Exit gate tổng", '=IF(AND(B11>=95%,B12=14,B13=0,B8=0,B9=0),"Đạt","Chưa đạt")', "Functional >=95%; Smoke 100%; không Blocker/Critical; không Blocked/Chưa chạy", ""],
    ]
    for row in metrics:
        ws.append(row)
    modules = ["Auth", "Course", "Class", "Enrollment", "Attendance", "Grade", "Common"]
    for offset, module in enumerate(modules, start=5):
        ws.cell(offset, 6, module)
        ws.cell(offset, 7, f'=COUNTIF(\'Test Cases\'!$B:$B,F{offset})')
        ws.cell(offset, 8, f'=COUNTIFS(\'Test Cases\'!$B:$B,F{offset},\'Test Cases\'!$L:$L,"Pass")&" / "&COUNTIFS(\'Test Cases\'!$B:$B,F{offset},\'Test Cases\'!$L:$L,"Fail")&" / "&COUNTIFS(\'Test Cases\'!$B:$B,F{offset},\'Test Cases\'!$L:$L,"Blocked")&" / "&COUNTIFS(\'Test Cases\'!$B:$B,F{offset},\'Test Cases\'!$L:$L,"Chưa chạy")')
    style_header(ws, 4, 1, 4)
    style_header(ws, 4, 6, 8)
    style_data(ws, 5, 15, 1, 4)
    style_data(ws, 5, 11, 6, 8)
    ws["B11"].number_format = "0.0%"
    ws["A17"] = "Module risk / nhận xét"
    ws["A17"].font = Font(bold=True, color=WHITE)
    ws["A17"].fill = PatternFill("solid", fgColor=BLUE)
    ws.merge_cells("B17:H17")
    ws["B17"] = "QA cập nhật sau vòng test: module rủi ro cao, bug còn mở, phạm vi chưa chạy."
    ws["A19"] = "QA sign-off"
    ws["B19"] = ""
    ws["D19"] = "Ngày"
    ws["E19"] = ""
    ws["A20"] = "Project sign-off"
    ws["B20"] = ""
    ws["D20"] = "Ngày"
    ws["E20"] = ""
    for row in ws.iter_rows(min_row=17, max_row=20, min_col=1, max_col=8):
        for cell in row:
            cell.border = Border(top=THIN_GRAY, bottom=THIN_GRAY, left=THIN_GRAY, right=THIN_GRAY)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 58
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 5
    ws.column_dimensions["F"].width = 16
    ws.column_dimensions["G"].width = 14
    ws.column_dimensions["H"].width = 34
    ws.sheet_view.showGridLines = False


def build_workbook():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.properties.title = "English Center MVC - QA Test Cases"
    wb.properties.subject = "Waterfall QA baseline and traceability"
    wb.properties.creator = "English Center Project Team"
    wb.properties.description = "Vietnamese QA workbook generated from the current MVC business rules."
    build_test_plan(wb)
    build_requirements(wb)
    build_test_cases(wb)
    build_rtm(wb)
    build_smoke(wb)
    build_regression(wb)
    build_bug_log(wb)
    build_summary(wb)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.save(OUTPUT)


def verify_workbook():
    wb = load_workbook(OUTPUT, data_only=False)
    expected_sheets = [
        "Test Plan", "Requirements", "Test Cases", "RTM", "Smoke Checklist",
        "Regression Checklist", "Bug Log", "Test Summary",
    ]
    assert wb.sheetnames == expected_sheets, wb.sheetnames
    ws = wb["Test Cases"]
    assert ws.max_row - 4 == 90, ws.max_row
    assert ws["A5"].value == "TC-AUTH-001"
    assert ws[f"A{ws.max_row}"].value == "TC-COMMON-006"
    known_requirements = {row[0] for row in REQUIREMENTS}
    used_requirements = {case["req"] for case in CASES}
    assert used_requirements == known_requirements, (known_requirements - used_requirements, used_requirements - known_requirements)
    assert len(wb["RTM"].tables) == 1
    assert len(wb["Bug Log"].data_validations.dataValidation) == 3
    return {
        "path": str(OUTPUT),
        "sheets": len(wb.sheetnames),
        "requirements": len(REQUIREMENTS),
        "test_cases": len(CASES),
        "smoke_items": wb["Smoke Checklist"].max_row - 4,
        "regression_items": wb["Regression Checklist"].max_row - 4,
    }


if __name__ == "__main__":
    build_workbook()
    result = verify_workbook()
    print("Workbook generated and verified:")
    for key, value in result.items():
        print(f"- {key}: {value}")
